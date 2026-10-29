import unittest
import os
import io
import sys
import json
import docx
import pypdf

ROOT_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if ROOT_DIR not in sys.path:
    sys.path.insert(0, ROOT_DIR)

from web_app import app
from routes.tts import parse_text_into_speech_segments


class TestTTSImportAndLongAudio(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.client = app.test_client()

    def test_parse_text_quotes_and_dialogues(self):
        """Kiểm tra tách câu với lời thoại trong ngoặc kép và văn bản dài."""
        dialogue = (
            '“Xin hỏi, đây có phải công ty của anh Vu Bân không ạ?” '
            '“Hiện tại chưa phải công ty. Nhưng nếu cô làm việc tốt, vài năm nữa có thể sẽ phải chuyển sang tòa nhà rộng hơn.”'
        )
        segments = parse_text_into_speech_segments(dialogue, default_pause=0.5)
        speech_segs = [s for s in segments if s[0] == 'speech']
        self.assertEqual(len(speech_segs), 3)
        self.assertIn("Vu Bân", speech_segs[0][1])
        self.assertIn("Hiện tại chưa phải công ty", speech_segs[1][1])
        self.assertIn("chuyển sang tòa nhà rộng hơn", speech_segs[2][1])

    def test_import_txt_file(self):
        """Kiểm tra import file .txt kịch bản."""
        content = "Đây là kịch bản truyện audio trích xuất từ tệp văn bản. Chương 1: Khởi đầu mới."
        data = {
            'file': (io.BytesIO(content.encode('utf-8')), 'truyen_ngan.txt')
        }
        res = self.client.post('/api/tts/import_file', data=data, content_type='multipart/form-data')
        self.assertEqual(res.status_code, 200)
        json_data = res.get_json()
        self.assertTrue(json_data.get('success'))
        self.assertIn("Chương 1", json_data.get('text'))
        self.assertEqual(json_data.get('suggested_title'), 'audio_truyen_ngan')
        self.assertGreater(json_data['stats']['word_count'], 10)

    def test_import_docx_file(self):
        """Kiểm tra import file Word .docx."""
        doc = docx.Document()
        doc.add_paragraph("Chương 1: Đêm mưa gió bão bùng tại căn nhà hoang.")
        doc.add_paragraph("Hắn cẩn thận bước từng bước vào trong bóng tối tĩnh mịch.")
        bio = io.BytesIO()
        doc.save(bio)
        bio.seek(0)

        data = {
            'file': (bio, 'truyen_ma.docx')
        }
        res = self.client.post('/api/tts/import_file', data=data, content_type='multipart/form-data')
        self.assertEqual(res.status_code, 200)
        json_data = res.get_json()
        self.assertTrue(json_data.get('success'))
        self.assertIn("Chương 1", json_data.get('text'))
        self.assertIn("bóng tối tĩnh mịch", json_data.get('text'))
        self.assertEqual(json_data.get('suggested_title'), 'audio_truyen_ma')

    def test_import_srt_file(self):
        """Kiểm tra import file phụ đề .srt bóc tách lời thoại."""
        srt_raw = (
            "1\n00:00:01,000 --> 00:00:03,500\nNgày xửa ngày xưa có một vương quốc.\n\n"
            "2\n00:00:04,000 --> 00:00:07,000\nNơi đó có một nàng công chúa xinh đẹp.\n"
        )
        data = {
            'file': (io.BytesIO(srt_raw.encode('utf-8')), 'kich_ban_phim.srt')
        }
        res = self.client.post('/api/tts/import_file', data=data, content_type='multipart/form-data')
        self.assertEqual(res.status_code, 200)
        json_data = res.get_json()
        self.assertTrue(json_data.get('success'))
        self.assertIn("Ngày xửa ngày xưa", json_data.get('text'))
        self.assertIn("nàng công chúa", json_data.get('text'))
        self.assertNotIn("-->", json_data.get('text'))

    def test_tts_generation_user_reported_bug_resolved(self):
        """Kiểm tra gọi TTS với đúng câu thoại bị lỗi ndarray.get trước đó."""
        payload = {
            "text": "“Xin hỏi, đây có phải công ty của anh Vu Bân không ạ?” “Hiện tại chưa phải công ty. Nhưng nếu cô làm việc tốt, vài năm nữa có thể sẽ phải chuyển sang tòa nhà rộng hơn.”",
            "voice": "local_clone_1787245769140",
            "speed": 1.0,
            "threads": 16,
            "output_dir": "output",
            "filename": "test_verification_minh_phuong.wav"
        }
        res = self.client.post('/api/tts/kokoro', json=payload)
        self.assertEqual(res.status_code, 200)
        data = res.get_json()
        self.assertTrue(data.get('success'))
        self.assertIn('audio_url', data)
        self.assertIn('srt_url', data)
        self.assertGreater(data.get('duration', 0), 5.0)

        # File tồn tại thực tế trên đĩa
        abs_audio = data.get('absolute_path')
        abs_srt = data.get('srt_absolute_path')
        self.assertTrue(os.path.exists(abs_audio))
        self.assertTrue(os.path.exists(abs_srt))
        self.assertGreater(os.path.getsize(abs_audio), 10000)


if __name__ == '__main__':
    unittest.main()
